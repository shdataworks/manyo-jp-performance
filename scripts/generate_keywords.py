import json
import sys
from datetime import date, datetime, timezone
from openpyxl import load_workbook

source, output = sys.argv[1:3]
workbook = load_workbook(source, read_only=True, data_only=True)
sheet = workbook['AMZ JP Keywords Raw']
rows = sheet.iter_rows(values_only=True)
source_headers = list(next(rows))
first_index = {}
for index, header in enumerate(source_headers):
    if header and header not in first_index:
        first_index[header] = index

headers = ['date', 'ad type', 'campaign', 'adset', 'asin', 'keyword', 'cost', 'imp', 'click', 'view', 'dpv', 'cart', 'purchase', 'sales', 'cart(new)', 'purchase(new)', 'sales(new)']
result = [headers]
for row in rows:
    values = []
    for header in headers:
        value = row[first_index[header]] if first_index[header] < len(row) else None
        if isinstance(value, (date, datetime)):
            value = value.strftime('%Y-%m-%d')
        values.append(value if value is not None else '')
    if any(value != '' for value in values):
        result.append(values)

payload = {
    'source': 'AMZ JP RAW',
    'sheet': 'AMZ JP Keywords Raw',
    'retrievedAt': datetime.now(timezone.utc).isoformat().replace('+00:00', 'Z'),
    'rows': result,
}
with open(output, 'w', encoding='utf-8') as file:
    json.dump(payload, file, ensure_ascii=False, separators=(',', ':'))
